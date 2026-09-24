"""
Calculation report for Lythos MSEW.

The report is assembled as HTML (inputs, earth pressure and forces, the
external checks and the bearing capacity by every method, the layer-by-layer
internal checks, the earthquake, the required length, the figures, the
warnings, the method notes and, if one was run, the height study) and
exported as
  * PDF   — through reportlab (see `lythosmsew.pdf`)
  * HTML  — a single self-contained file (figures embedded as base64)
  * DOCX  — optional, needs `python-docx`

There is one assembly, `build_html()`, so all three formats say the same thing.
"""

from __future__ import annotations

import base64
import datetime
import html
import io
import math
from typing import Any, Dict, List, Optional

from matplotlib.figure import Figure

from . import height_plots
from . import reinforcement as R
from .config import APP_NAME, APP_VERSION
from .i18n import TRANSLATIONS, warning_text
from .plotting import Plotter, available_figures
from .summary import EXTERNAL, INTERNAL, bearing_table, value_name

TEXTS = {
    "en": {
        "title": f"{APP_NAME} — Calculation Report",
        "date": "Date", "analyst": "Analyst", "software": "Software",
        "sec_inputs": "1. Input data", "sec_geom": "1.1 Wall geometry and surcharges",
        "sec_soils": "1.2 Soils", "sec_types": "1.3 Reinforcement types",
        "sec_layers": "1.4 Reinforcement layers", "sec_opts": "1.5 Design method and criteria",
        "sec_forces": "2. Earth pressure and forces",
        "sec_ka": "2.1 Earth pressure coefficients", "sec_force_list": "2.2 Forces on the "
        "reinforced block (per metre)",
        "sec_external": "3. External stability", "sec_ext_checks": "3.1 Checks",
        "sec_bearing": "3.2 Bearing capacity of the foundation",
        "sec_factors": "3.3 Factors of the chosen method",
        "sec_internal": "4. Internal stability", "sec_tension": "4.1 Tension in the layers",
        "sec_pullout": "4.2 Pullout", "sec_int_checks": "4.3 Checks, layer by layer",
        "sec_seismic": "5. Earthquake (pseudo-static)", "sec_seis_ext": "5.1 External stability",
        "sec_seis_int": "5.2 Internal stability",
        "sec_length": "6. The reinforcement length the checks need",
        "sec_figs": "7. Figures", "sec_warn": "8. Warnings", "sec_notes": "9. Method notes",
        "sec_heights": "10. Height study", "sec_heights_table": "10.1 Results by height",
        "sec_heights_figs": "10.2 Figures",
        "parameter": "Parameter", "value": "Value", "unit": "Unit",
        "H": "Design height H", "d": "Embedment d", "omega": "Face batter ω",
        "beta": "Backslope β", "toe": "Slope in front of the toe",
        "q_dead": "Permanent surcharge", "q_live": "Live surcharge",
        "strip": "Strip load P (width, face to centre)", "live": "live", "dead": "permanent",
        "soil": "Soil", "gamma": "γ kN/m³", "gamma_sat": "γsat kN/m³", "phi": "φ' °",
        "c": "c' kPa", "zw": "Water table below the base", "gw": "γw",
        "name": "Name", "kind": "Kind", "strength": "T_al kN/m (per m of wall)",
        "strength_dyn": "T for seismic kN/m", "coverage": "Rc", "detail": "Detail",
        "z": "z m", "L": "L m", "type": "Type",
        "design": "Design method", "bearing_method": "Bearing capacity factors",
        "embedment": "Embedment as surcharge", "inclination": "Load inclination factors",
        "Cds": "Direct sliding coefficient Cds", "life": "Design life (steel)",
        "yes": "yes", "no": "no",
        "Ka_b": "Ka of the retained fill (δ = β)", "Ka_r": "Ka of the reinforced fill",
        "theta": "Inclination of the back θ", "sigma2": "Backslope as a surcharge σ2",
        "force": "Force", "magnitude": "kN/m", "arm": "Arm m", "moment": "kNm/m",
        "V": "ΣV", "Hf": "ΣH", "MR": "Resisting moment M_R", "MO": "Overturning moment M_O",
        "e": "Eccentricity e", "check": "Check", "required": "Required", "status": "Status",
        "B_eff": "Effective width B' = L − 2e", "sigma_v": "Vertical stress σv = ΣV / B'",
        "q_ult": "Ultimate capacity q_ult", "factor": "Factor",
        "c_term": "cohesion", "q_term": "surcharge", "g_term": "self weight",
        "ok": "OK", "notok": "NOT OK", "na": "n/a",
        "length_found": "The shortest uniform length that satisfies the external and pullout "
                        "checks is L = {L:.2f} m; FHWA's minimum is {rule:.2f} m; the wall has "
                        "L = {cur:.2f} m.",
        "length_none": "No uniform length up to 4·H satisfies the checks.",
        "notes": [
            "External stability follows FHWA-NHI-10-024 and AASHTO LRFD 11.10.5. The "
            "reinforced block is a rigid body whose length is that of its lowest layer; the "
            "retained fill pushes on its back with Coulomb's Ka, the thrust inclined at the "
            "backslope angle, and the surcharges on the retained fill add Ka·q·h. The live "
            "load over the reinforced zone is left out of sliding and eccentricity and kept "
            "in bearing.",
            "Sliding resists with the weakest of the reinforced fill (tan φr), the foundation "
            "soil (tan φf and c·L) and, under a geosynthetic lowest layer, Cds·tan φr.",
            "The foundation is a strip footing of width L. The resultant's eccentricity leaves "
            "Meyerhof's effective width B' = L − 2e and the vertical stress ΣV / B', compared "
            "with q_ult of the general bearing capacity equation (the embedment and the load "
            "inclination only when switched on, as FHWA leaves them out; a slope in front of "
            "the toe through the ground factors).",
            "Internal stability follows AASHTO's Simplified Method. The tension in a layer is "
            "Tmax = Kr·σv·Sv with Sv its tributary height; Kr/Ka is 1 for geosynthetics, and "
            "for steel strips falls from 1.7 at the top to 1.2 at 6 m. σv holds the fill, the "
            "backslope as the surcharge σ2 = ½·(0.7H)·tan β·γr, the surcharges and the 2:1 "
            "spread of a strip load.",
            "The active zone is Rankine's plane at 45 + φ/2 for extensible reinforcement and "
            "the bilinear 0.3·H1 line for inextensible reinforcement. Pullout resistance "
            "Pr = F*·α·σ'v·Le·C·Rc with C = 2, σ'v without live load; F* is Ci·tan φ for "
            "geosynthetics and falls from F*₀ at the top to tan φ at 6 m for ribbed strips.",
            "A geosynthetic's long-term strength is Tult / (RFID·RFCR·RFD) times its coverage "
            "ratio; a steel strip's is Fy·b·Ec per strip divided by its spacing, Ec the "
            "thickness left after the zinc (15 µm/yr for 2 years, then 4 µm/yr) and then the "
            "steel on both faces are lost over the design life. In ASD the steel may carry "
            "0.55·Fy·Ac and a geosynthetic T_al / FS; in LRFD φ·T_al.",
            "ASD asks for a factor of safety on every check. LRFD factors the loads (EV 1.00 "
            "or 1.35, EH 1.50, ES 0.75 or 1.50, LS 1.75) and the resistances (φ) and asks for "
            "a capacity-to-demand ratio of at least one.",
            "The required length is the shortest uniform length that satisfies the external "
            "checks and the pullout of every layer, static and seismic, found by bisection.",
            "Not included: global and compound stability, settlement, and the drainage of the "
            "wall; check them separately.",
        ],
        "note_seismic": "Earthquake (FHWA-NHI-10-024 §7): Am = (1.45 − A)·A; the dynamic "
                        "thrust PAE = 0.375·Am·γ·H² (Mononobe–Okabe with a backslope) acts at "
                        "0.6·H, half of it together with the inertia PIR of a block 0.5·H wide; "
                        "internally the inertia of the active zone Pi = Am·Wa is shared out by "
                        "the resisting lengths, the geosynthetic's dynamic part carried by "
                        "Tult / (RFID·RFD), F* reduced by the given factor, and the live load "
                        "halved.",
    },
    "tr": {
        "title": f"{APP_NAME} — Hesap Raporu",
        "date": "Tarih", "analyst": "Hazırlayan", "software": "Yazılım",
        "sec_inputs": "1. Girdi verileri", "sec_geom": "1.1 Duvar geometrisi ve sürşarjlar",
        "sec_soils": "1.2 Zeminler", "sec_types": "1.3 Donatı türleri",
        "sec_layers": "1.4 Donatı tabakaları", "sec_opts": "1.5 Tasarım yöntemi ve ölçütler",
        "sec_forces": "2. Toprak basıncı ve kuvvetler",
        "sec_ka": "2.1 Toprak basıncı katsayıları",
        "sec_force_list": "2.2 Donatılı bloğa etkiyen kuvvetler (metre başına)",
        "sec_external": "3. Dış duraylılık", "sec_ext_checks": "3.1 Kontroller",
        "sec_bearing": "3.2 Temel zemininin taşıma gücü",
        "sec_factors": "3.3 Seçilen yöntemin katsayıları",
        "sec_internal": "4. İç duraylılık", "sec_tension": "4.1 Tabakalardaki çekme kuvveti",
        "sec_pullout": "4.2 Sıyrılma", "sec_int_checks": "4.3 Kontroller, tabaka tabaka",
        "sec_seismic": "5. Deprem (psödo-statik)", "sec_seis_ext": "5.1 Dış duraylılık",
        "sec_seis_int": "5.2 İç duraylılık",
        "sec_length": "6. Kontrollerin gerektirdiği donatı boyu",
        "sec_figs": "7. Şekiller", "sec_warn": "8. Uyarılar", "sec_notes": "9. Yöntem notları",
        "sec_heights": "10. Yükseklik çalışması",
        "sec_heights_table": "10.1 Yüksekliğe göre sonuçlar", "sec_heights_figs": "10.2 Şekiller",
        "parameter": "Parametre", "value": "Değer", "unit": "Birim",
        "H": "Tasarım yüksekliği H", "d": "Gömülme derinliği d", "omega": "Yüz eğimi ω",
        "beta": "Arka şev β", "toe": "Topuk önündeki şev",
        "q_dead": "Kalıcı sürşarj", "q_live": "Hareketli sürşarj",
        "strip": "Şerit yük P (genişlik, yüzden merkeze)", "live": "hareketli",
        "dead": "kalıcı",
        "soil": "Zemin", "gamma": "γ kN/m³", "gamma_sat": "γdoy kN/m³", "phi": "φ' °",
        "c": "c' kPa", "zw": "Su tablasının tabandan derinliği", "gw": "γw",
        "name": "Ad", "kind": "Tür", "strength": "T_al kN/m (duvarın metresi başına)",
        "strength_dyn": "Deprem için T kN/m", "coverage": "Rc", "detail": "Ayrıntı",
        "z": "z m", "L": "L m", "type": "Tür",
        "design": "Tasarım yöntemi", "bearing_method": "Taşıma gücü katsayıları",
        "embedment": "Gömülme sürşarj olarak", "inclination": "Yük eğim katsayıları",
        "Cds": "Doğrudan kayma katsayısı Cds", "life": "Tasarım ömrü (çelik)",
        "yes": "evet", "no": "hayır",
        "Ka_b": "Arka dolgunun Ka'sı (δ = β)", "Ka_r": "Donatılı dolgunun Ka'sı",
        "theta": "Arka yüzün eğimi θ", "sigma2": "Arka şevin sürşarj karşılığı σ2",
        "force": "Kuvvet", "magnitude": "kN/m", "arm": "Kol m", "moment": "kNm/m",
        "V": "ΣV", "Hf": "ΣH", "MR": "Karşı koyan moment M_R", "MO": "Devirici moment M_O",
        "e": "Dışmerkezlik e", "check": "Kontrol", "required": "Gerekli", "status": "Durum",
        "B_eff": "Etkin genişlik B' = L − 2e", "sigma_v": "Düşey gerilme σv = ΣV / B'",
        "q_ult": "Nihai taşıma gücü q_ult", "factor": "Katsayı",
        "c_term": "kohezyon", "q_term": "sürşarj", "g_term": "zati ağırlık",
        "ok": "UYGUN", "notok": "UYGUN DEĞİL", "na": "—",
        "length_found": "Dış duraylılık ve sıyrılma kontrollerini sağlayan en kısa eşit boy "
                        "L = {L:.2f} m; FHWA'nın en küçüğü {rule:.2f} m; duvarda "
                        "L = {cur:.2f} m.",
        "length_none": "4·H'ye kadar hiçbir eşit boy kontrolleri sağlamıyor.",
        "notes": [
            "Dış duraylılık FHWA-NHI-10-024 ve AASHTO LRFD 11.10.5'e göredir. Donatılı blok, "
            "boyu en alt tabakanın boyu olan rijit bir cisimdir; arka dolgu arka yüzüne "
            "Coulomb Ka ile, arka şev açısıyla eğik bir itki uygular ve arka dolgu üzerindeki "
            "sürşarjlar Ka·q·h ekler. Donatılı bölge üzerindeki hareketli yük kayma ve "
            "dışmerkezlikte dikkate alınmaz, taşıma gücünde alınır.",
            "Kayma direnci; donatılı dolgu (tan φr), temel zemini (tan φf ve c·L) ve en alt "
            "tabaka geosentetik ise Cds·tan φr değerlerinin en zayıfıdır.",
            "Temel, L genişliğinde şerit temeldir. Bileşkenin dışmerkezliği Meyerhof'un etkin "
            "genişliğini B' = L − 2e ve düşey gerilmeyi ΣV / B' verir; bu gerilme genel taşıma "
            "gücü denkleminin q_ult değeriyle karşılaştırılır (gömülme ve yük eğimi, FHWA "
            "dikkate almadığı için yalnızca açıldığında; topuk önündeki şev zemin eğim "
            "katsayılarıyla).",
            "İç duraylılık AASHTO Basitleştirilmiş Yöntemi'ne göredir. Tabakadaki çekme "
            "kuvveti Tmax = Kr·σv·Sv, Sv tabakanın etki yüksekliğidir; Kr/Ka geosentetikte 1, "
            "çelik şeritte tepede 1.7'den 6 m'de 1.2'ye azalır. σv; dolguyu, σ2 = "
            "½·(0.7H)·tan β·γr sürşarjı olarak arka şevi, sürşarjları ve şerit yükün 2:1 "
            "yayılışını içerir.",
            "Aktif bölge, uzayabilir donatıda 45 + φ/2 açılı Rankine düzlemi, uzamaz donatıda "
            "iki doğrulu 0.3·H1 çizgisidir. Sıyrılma direnci Pr = F*·α·σ'v·Le·C·Rc, C = 2 ve "
            "σ'v hareketli yük olmadan; F* geosentetikte Ci·tan φ, nervürlü şeritte tepede "
            "F*₀'dan 6 m'de tan φ'ye azalır.",
            "Geosentetiğin uzun süreli dayanımı Tult / (RFID·RFCR·RFD) çarpı kaplama oranıdır; "
            "çelik şeridinki şerit başına Fy·b·Ec bölü aralıktır. Ec, tasarım ömrü boyunca "
            "çinko (2 yıl 15 µm/yıl, sonra 4 µm/yıl) ve ardından iki yüzden çelik kaybedildikten "
            "sonra kalan kalınlıktır. ASD'de çelik 0.55·Fy·Ac, geosentetik T_al / GS taşır; "
            "LRFD'de φ·T_al.",
            "ASD her kontrol için güvenlik sayısı ister. LRFD yükleri (EV 1.00 veya 1.35, "
            "EH 1.50, ES 0.75 veya 1.50, LS 1.75) ve dirençleri (φ) çarpanlar ve en az bir "
            "kapasite/talep oranı (KTO) ister.",
            "Gerekli boy, dış duraylılık kontrollerini ve her tabakanın sıyrılmasını, statik "
            "ve depremde, sağlayan en kısa eşit boydur ve ikiye bölme ile bulunur.",
            "Kapsam dışı: genel ve bileşik duraylılık, oturma ve duvar drenajı; ayrıca "
            "kontrol ediniz.",
        ],
        "note_seismic": "Deprem (FHWA-NHI-10-024 §7): Am = (1.45 − A)·A; dinamik itki "
                        "PAE = 0.375·Am·γ·H² (arka şevde Mononobe–Okabe) 0.6·H'de etkir, "
                        "yarısı 0.5·H genişliğindeki bloğun ataletiyle (PIR) birlikte alınır; "
                        "iç duraylılıkta aktif bölgenin ataleti Pi = Am·Wa direnç boylarına göre "
                        "paylaştırılır, geosentetiğin dinamik payını Tult / (RFID·RFD) taşır, "
                        "F* verilen katsayıyla azaltılır ve hareketli yük yarıya indirilir.",
    },
}

#: The interface's palette: warm paper, ink, terracotta; serif headings
_CSS = """
body { font-family: system-ui, -apple-system, 'Segoe UI', 'Helvetica Neue', Arial, sans-serif;
       font-size: 10pt; color: #141413; }
h1, h2, h3 { font-family: 'Tiempos Text', 'Source Serif 4', 'Iowan Old Style', Palatino,
             Georgia, 'DejaVu Serif', serif; font-weight: 500; }
h1 { font-size: 20pt; color: #141413; margin-bottom: 2px; }
h2 { font-size: 14pt; color: #C6613F; border-bottom: 1px solid #E3E0D5; padding-bottom: 2px;
     margin-top: 20px; }
h3 { font-size: 11.5pt; color: #141413; margin-top: 12px; }
table { border-collapse: collapse; margin: 4px 0 8px 0; }
th { background: #F0EEE6; text-align: left; padding: 3px 6px; border: 1px solid #E3E0D5;
     font-size: 9pt; }
td { padding: 3px 6px; border: 1px solid #E3E0D5; font-size: 9pt; }
.ok { color: #3F7F4F; font-weight: bold; } .bad { color: #B0413E; font-weight: bold; }
.meta { color: #73726C; } .note { color: #73726C; font-size: 9pt; }
"""


def _esc(x) -> str:
    return html.escape(str(x))


def _f(x, nd=2) -> str:
    if x is None or (isinstance(x, float) and not math.isfinite(x)):
        return "∞" if x == math.inf else "—"
    return f"{x:,.{nd}f}"


def _status(T, s: str) -> str:
    if s == "N/A":
        return T["na"]
    return (f'<span class="ok">{T["ok"]}</span>' if s == "OK"
            else f'<span class="bad">{T["notok"]}</span>')


def _table(headers: List[str], rows: List[List[Any]], widths: Optional[List[int]] = None) -> str:
    out = ["<table width='100%'>"]
    if widths:
        cells = "".join(f"<th width='{w}%'>{_esc(h)}</th>" for h, w in zip(headers, widths))
    else:
        cells = "".join(f"<th>{_esc(h)}</th>" for h in headers)
    out.append("<tr>" + cells + "</tr>")
    for r in rows:
        out.append("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def _kv_table(T, rows: List[List[Any]]) -> str:
    return _table([T["parameter"], T["value"], T["unit"]], rows, [50, 36, 14])


def _check_table(T, L, design: str, checks: dict, names: List[str], internal=False) -> str:
    rows = []
    for key in names:
        check = checks[key]
        label = L["chk_internal_sliding"] if (internal and key == "sliding") else L[f"chk_{key}"]
        if internal and math.isfinite(check.get("z", math.nan)):
            label += f" (z = {check['z']:.3f} m)"
        nd = 3 if check["kind"] == "limit" else 2
        rows.append([_esc(label), _f(check["value"], nd), _f(check["required"], nd),
                     _status(T, check["status"])])
    return _table([T["check"], value_name(L, design), T["required"], T["status"]], rows,
                  [46, 18, 18, 18])


# ----------------------------------------------------------------------
def _png(fig: Figure, dpi: int) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
    return buf.getvalue()


def render_figures(wall, lang: str, dpi: int = 130) -> Dict[str, bytes]:
    """The report figures as PNG bytes (off-screen, white, untitled)."""
    plotter = Plotter(wall, lang, "paper", titles=False)
    out = {}
    for key in available_figures(wall):
        fig = Figure(figsize=(10, 7), dpi=dpi)
        plotter.draw(key, fig)
        out[key] = _png(fig, dpi)
    return out


def render_height_figures(study, lang: str, dpi: int = 110) -> Dict[str, bytes]:
    """Height study figures keyed 'hs_<view>'."""
    if study is None or not study.rows:
        return {}
    from .render import height_figure
    return {f"hs_{view}": _png(height_figure(study, view, lang, "paper", titles=False), dpi)
            for view in height_plots.VIEWS}


def _heights_section(T, L, study, figures, img_src) -> List[str]:
    parts = [f"<h2 style='page-break-before:always'>{T['sec_heights']}</h2>"]
    parts.append(f"<p>{_esc(height_plots.verdict(study, L))}</p>")
    parts.append(f"<h3>{T['sec_heights_table']}</h3>")
    tab = height_plots.table(study, L)
    rows = []
    for row in tab["rows"]:
        cells = []
        for cell, state in zip(row["cells"], row["states"]):
            cells.append(f'<span class="{state}">{_esc(cell)}</span>' if state else _esc(cell))
        rows.append(cells)
    parts.append(_table(tab["columns"], rows))
    parts.append(f"<h3>{T['sec_heights_figs']}</h3>")
    for view in height_plots.VIEWS:
        key = f"hs_{view}"
        if key in figures:
            parts.append(f"<p class='lead'>{_esc(L['fig_hs_' + view])}</p>")
            parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")
    return parts


def build_html(wall, lang: str, figures: Dict[str, bytes], img_src=None, heights=None) -> str:
    """
    Builds the report HTML. img_src(key) -> value for the <img src> attribute;
    the default embeds base64 data URIs (the self-contained HTML file). The PDF
    and DOCX writers pass a `fig://<key>` mapper and resolve the keys against
    the figure dictionary themselves.
    """
    T = TEXTS.get(lang, TEXTS["en"])
    L = TRANSLATIONS.get(lang, TRANSLATIONS["en"])
    w, res = wall, wall.results
    design = res["design"]
    lrfd = design == "lrfd"
    info = w.config.get("project_info", {})
    if img_src is None:
        def img_src(key):
            return "data:image/png;base64," + base64.b64encode(figures[key]).decode("ascii")

    parts = [f"<html><head><meta charset='utf-8'><style>{_CSS}</style></head><body>"]
    parts.append(f"<h1>{_esc(info.get('title') or T['title'])}</h1>")
    stamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    meta = f"{T['date']}: {stamp}"
    if info.get("analyst"):
        meta += f" &nbsp;·&nbsp; {T['analyst']}: {_esc(info['analyst'])}"
    meta += f" &nbsp;·&nbsp; {T['software']}: {APP_NAME} v{APP_VERSION}"
    parts.append(f"<p class='meta'>{meta}</p>")

    # ---------------------------------------------------------------- inputs
    parts.append(f"<h2>{T['sec_inputs']}</h2>")
    parts.append(f"<h3>{T['sec_geom']}</h3>")
    rows = [[T["H"], _f(w.H), "m"], [T["d"], _f(w.d), "m"], [T["omega"], _f(w.omega, 1), "°"],
            [T["beta"], _f(w.beta, 1), "°"], [T["toe"], _f(w.toe_slope, 1), "°"],
            [T["q_dead"], _f(w.q_dead, 1), "kPa"], [T["q_live"], _f(w.q_live, 1), "kPa"]]
    if w.strip_on:
        rows.append([T["strip"], f"{_f(w.strip_P, 1)} ({_f(w.strip_b)} m, {_f(w.strip_x)} m) — "
                     f"{T['live'] if w.strip_live else T['dead']}", "kN/m"])
    parts.append(_kv_table(T, rows))

    parts.append(f"<h3>{T['sec_soils']}</h3>")
    rows = [[L["soil_reinforced"], _f(w.rf["gamma"], 1), "—", _f(w.rf["phi"], 1), "—"],
            [L["soil_retained"], _f(w.rb["gamma"], 1), "—", _f(w.rb["phi"], 1), "—"],
            [L["soil_foundation"], _f(w.fd["gamma"], 1), _f(w.fd["gamma_sat"], 1),
             _f(w.fd["phi"], 1), _f(w.fd["c"], 1)]]
    parts.append(_table([T["soil"], T["gamma"], T["gamma_sat"], T["phi"], T["c"]], rows,
                        [32, 17, 17, 17, 17]))
    parts.append(f"<p>{T['zw']}: {_f(w.dw)} m &nbsp;·&nbsp; {T['gw']} = {_f(w.gw)} kN/m³</p>")

    parts.append(f"<h3>{T['sec_types']}</h3>")
    used = {layer["type"] for layer in w.layers}
    rows = []
    for name, t in w.types.items():
        if name not in used:
            continue
        st = R.strength(t, w.corrosion)
        if R.is_extensible(t["kind"]):
            detail = (f"Tult = {_f(t['Tult'], 1)} kN/m, RF = {_f(t['RFID'])}·{_f(t['RFCR'])}·"
                      f"{_f(t['RFD'])}, Ci = {_f(t['Ci'])}, α = {_f(t['alpha'])}, "
                      f"CR = {_f(t['CR'])}")
        else:
            detail = (f"b × t = {_f(t['b'], 0)} × {_f(t['t'])} mm, Fy = {_f(t['Fy'], 0)} MPa, "
                      f"Sh = {_f(t['Sh'])} m, Es = {_f(st['Es'], 3)} mm, "
                      f"Ec = {_f(st['Ec'], 3)} mm, F*₀ = {_f(t['F0'])}, α = {_f(t['alpha'])}")
        rows.append([_esc(name), L[f"kind_{t['kind']}"], _esc(detail), _f(st["T_lt"], 1),
                     _f(st["Rc"], 3)])
    parts.append(_table([T["name"], T["kind"], T["detail"], T["strength"], T["coverage"]],
                        rows, [16, 12, 48, 14, 10]))

    parts.append(f"<h3>{T['sec_layers']}</h3>")
    rows = [[_f(r["z"], 3), _f(r["L"]), _esc(r["type"])] for r in res["layers"]]
    parts.append(_table([T["z"], T["L"], T["type"]], rows, [30, 30, 40]))

    parts.append(f"<h3>{T['sec_opts']}</h3>")
    crit = w.crit
    rows = [[T["design"], L[f"design_{design}"], ""],
            [T["bearing_method"], L[f"method_{w.method}"], ""],
            [T["embedment"], T["yes"] if w.use_embedment else T["no"], ""],
            [T["inclination"], T["yes"] if w.use_inclination else T["no"], ""],
            [T["Cds"], _f(w.Cds), ""],
            [T["life"], _f(w.corrosion["design_life"], 0), "yr"]]
    if lrfd:
        for key in ("phi_sliding", "phi_bearing", "phi_steel", "phi_geo", "phi_pullout"):
            rows.append([L[f"{key}_label"], _f(crit[key]), ""])
        rows.append([L["ecc_lrfd_label"], _f(crit["ecc_lrfd"], 1), ""])
    else:
        for key in ("FS_sliding", "FS_overturning", "FS_bearing", "FS_tensile", "FS_pullout",
                    "FS_connection"):
            rows.append([f"FS — {L[key + '_label']}", _f(crit[key]), ""])
        rows.append([L["steel_ratio_label"], _f(crit["steel_ratio"]), ""])
        rows.append([L["ecc_asd_label"], _f(crit["ecc_asd"], 1), ""])
    rows.append([L["min_Le_label"], _f(crit["min_Le"]), "m"])
    parts.append(_kv_table(T, rows))

    # ---------------------------------------------------------- forces
    parts.append(f"<h2>{T['sec_forces']}</h2>")
    parts.append(f"<h3>{T['sec_ka']}</h3>")
    parts.append(_kv_table(T, [
        [T["Ka_b"], _f(res["Ka"]["retained"], 4), ""],
        [T["Ka_r"], _f(res["Ka"]["reinforced"], 4), ""],
        [T["theta"], _f(res["Ka"]["theta"], 1), "°"],
        [T["sigma2"], _f(res["slope_surcharge"], 2), "kPa"],
    ]))
    parts.append(f"<h3>{T['sec_force_list']}</h3>")
    f = res["forces"]
    rows = [[_esc(n), _f(v, 1), _f(x, 3), _f(v * x, 1)] for n, v, x in f["vertical"]]
    rows += [[_esc(n), _f(h, 1), _f(y, 3), _f(h * y, 1)] for n, h, y in f["horizontal"]]
    parts.append(_table([T["force"], T["magnitude"], T["arm"], T["moment"]], rows,
                        [34, 22, 22, 22]))
    parts.append(_kv_table(T, [
        [T["V"], _f(f["V"], 1), "kN/m"], [T["Hf"], _f(f["H"], 1), "kN/m"],
        [T["MR"], _f(f["M_R"], 1), "kNm/m"], [T["MO"], _f(f["M_O"], 1), "kNm/m"],
        [T["e"], _f(f["e"], 3), "m"],
    ]))

    # ---------------------------------------------------------- external
    parts.append(f"<h2>{T['sec_external']}</h2>")
    parts.append(f"<h3>{T['sec_ext_checks']}</h3>")
    parts.append(_check_table(T, L, design, res["external"], EXTERNAL))
    plane = res["external"]["sliding"].get("plane", "reinforced")
    sentence = L["res_sliding_plane"].format(plane=L["plane_" + plane]).strip()
    parts.append(f"<p>{_esc(sentence[:1].upper() + sentence[1:])}</p>")

    b = res["bearing"]
    parts.append(f"<h3>{T['sec_bearing']}</h3>")
    parts.append(_kv_table(T, [
        [T["B_eff"], _f(b["B_eff"], 3), "m"], [T["sigma_v"], _f(b.get("sigma_v"), 1), "kPa"],
        [T["q_ult"] + f" ({L['method_' + b['primary']]})", _f(b.get("q_ult"), 1), "kPa"],
    ]))
    tab = bearing_table(w, lang)
    rows = []
    for row in tab["rows"]:
        cells = [_esc(c) for c in row["cells"]]
        if row.get("primary"):
            cells[0] = f"<b>{cells[0]} *</b>"
        cells[-1] = f'<span class="{row["states"][-1]}">{cells[-1]}</span>'
        rows.append(cells)
    parts.append(_table(tab["columns"], rows, [26, 11, 11, 11, 14, 14, 13]))
    primary = b["methods"].get(b["primary"])
    if primary:
        parts.append(f"<h3>{T['sec_factors']}</h3>")
        frows = [[L.get(f"factor_{name}", name), _f(primary["factors"][name]["c"], 3),
                  _f(primary["factors"][name]["q"], 3), _f(primary["factors"][name]["g"], 3)]
                 for name in ("inclination", "ground")]
        frows.append(["N", _f(primary["N"]["Nc"], 2), _f(primary["N"]["Nq"], 2),
                      _f(primary["N"]["Ngamma"], 2)])
        frows.append(["kPa", _f(primary["terms"]["c"], 1), _f(primary["terms"]["q"], 1),
                      _f(primary["terms"]["g"], 1)])
        parts.append(_table([T["factor"], T["c_term"], T["q_term"], T["g_term"]], frows,
                            [40, 20, 20, 20]))

    # ---------------------------------------------------------- internal
    parts.append(f"<h2>{T['sec_internal']}</h2>")
    parts.append(f"<h3>{T['sec_tension']}</h3>")
    rows = [[_f(r["z"], 3), _esc(r["type"]), _f(r["Sv"], 3), _f(r["Kr"], 3),
             _f(r["sigma_v"], 1), _f(r["T_design"] if lrfd else r["T_max"]),
             _f(r["T_lt"], 1), _f(r["tensile"]["capacity"], 1)]
            for r in res["layers"]]
    parts.append(_table([L["head_z"], L["head_type"], L["head_Sv"], L["head_Kr"],
                         L["head_sigma_v"], L["head_Tmax"], L["head_Tlt"],
                         "φ·T_al" if lrfd else "T_a"], rows, [10, 18, 10, 10, 12, 14, 13, 13]))
    parts.append(f"<h3>{T['sec_pullout']}</h3>")
    rows = [[_f(r["z"], 3), _f(r["L"]), _f(r["La"]), _f(r["Le"]), _f(r["sigma_p"], 1),
             _f(r["F_star"], 3), _f(r["alpha"]), _f(r["Rc"], 3), _f(r["P_r"], 1)]
            for r in res["layers"]]
    parts.append(_table([L["head_z"], L["col_L"], L["head_La"], L["head_Le"], "σ'v kPa",
                         L["head_Fstar"], "α", "Rc", L["head_Pr"]], rows,
                        [10, 10, 10, 10, 12, 11, 9, 10, 18]))
    parts.append(f"<h3>{T['sec_int_checks']}</h3>")
    rows = []
    for r in res["layers"]:
        cells = [_f(r["z"], 3)]
        for key in INTERNAL:
            chk = r[key]
            cells.append(f"{_f(chk['value'])} / {_f(chk['required'])} "
                         f"{_status(T, chk['status'])}")
        rows.append(cells)
    parts.append(_table([L["head_z"], L["head_tensile"], L["head_pullout"],
                         L["head_connection"], L["head_sliding"]], rows, [12, 22, 22, 22, 22]))
    parts.append(_check_table(T, L, design, res["internal"], INTERNAL, internal=True))

    # ---------------------------------------------------------- seismic
    s = res["seismic"]
    if s:
        parts.append(f"<h2>{T['sec_seismic']}</h2>")
        parts.append(f"<p>{_esc(L['res_seismic'].format(A=w.A, Am=s['Am'], PAE=s['PAE'], PIR=s['PIR'], Pi=s['P_i']))}</p>")
        parts.append(f"<h3>{T['sec_seis_ext']}</h3>")
        parts.append(_check_table(T, L, design, s["external"], EXTERNAL))
        parts.append(f"<h3>{T['sec_seis_int']}</h3>")
        rows = [[_f(r["z"], 3), _f(r["T_static"]), _f(r["T_md"]), _f(r["T_total"]),
                 _f(r["P_r"], 1),
                 f"{_f(r['tensile']['value'])} {_status(T, r['tensile']['status'])}",
                 f"{_f(r['pullout']['value'])} {_status(T, r['pullout']['status'])}"]
                for r in s["internal_rows"]]
        parts.append(_table([L["head_z"], L["head_Tmax"], L["head_Tmd"], L["head_Ttotal"],
                             L["head_Pr"], L["head_tensile"], L["head_pullout"]], rows,
                            [10, 13, 13, 13, 13, 19, 19]))
        parts.append(_check_table(T, L, design, s["internal"],
                                  ["tensile", "pullout", "connection"], internal=True))

    # ---------------------------------------------------------- length
    parts.append(f"<h2>{T['sec_length']}</h2>")
    length = res["required_length"]
    text = (T["length_found"].format(L=length, rule=res["length_rule"], cur=res["L"])
            if math.isfinite(length) else T["length_none"])
    parts.append(f"<p>{_esc(text)}</p>")

    # ---------------------------------------------------------- figures
    parts.append(f"<h2 style='page-break-before:always'>{T['sec_figs']}</h2>")
    for key in available_figures(w):
        if key not in figures:
            continue
        parts.append(f"<p class='lead'>{_esc(L.get('fig_' + key, key))}</p>")
        parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")

    if res["warnings"]:
        parts.append(f"<h2>{T['sec_warn']}</h2><ul>")
        parts.extend(f"<li>{_esc(warning_text(lang, x))}</li>" for x in res["warnings"])
        parts.append("</ul>")

    parts.append(f"<h2>{T['sec_notes']}</h2><ul class='note'>")
    notes = list(T["notes"])
    if s:
        notes.append(T["note_seismic"])
    parts.extend(f"<li>{_esc(n)}</li>" for n in notes)
    parts.append("</ul>")

    if heights is not None and heights.rows:
        parts.extend(_heights_section(T, L, heights, figures, img_src))
    parts.append("</body></html>")
    return "\n".join(parts)


# ----------------------------------------------------------------------
# exporters
# ----------------------------------------------------------------------
def _all_figures(wall, lang, heights):
    figures = render_figures(wall, lang)
    figures.update(render_height_figures(heights, lang))
    return figures


def export_html(path: str, wall, lang: str, heights=None) -> None:
    figures = _all_figures(wall, lang, heights)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(build_html(wall, lang, figures, heights=heights))


def export_pdf(path: str, wall, lang: str, heights=None) -> None:
    """PDF through reportlab; the figures travel as `fig://<key>` references."""
    from . import pdf as pdf_writer

    T = TEXTS.get(lang, TEXTS["en"])
    info = wall.config.get("project_info", {})
    figures = _all_figures(wall, lang, heights)
    html_text = build_html(wall, lang, figures, img_src=lambda k: f"fig://{k}", heights=heights)
    pdf_writer.html_to_pdf(path, html_text, figures,
                           title=info.get("title", T["title"]),
                           author=info.get("analyst", ""),
                           footer=f"{info.get('title', '')} — {APP_NAME} v{APP_VERSION}")


def export_docx(path: str, wall, lang: str, heights=None) -> None:
    """Word report; requires python-docx."""
    try:
        import docx
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed (pip install python-docx).") from exc
    import re

    figures = _all_figures(wall, lang, heights)
    html_text = build_html(wall, lang, figures, img_src=lambda k: f"fig://{k}", heights=heights)
    d = docx.Document()
    from docx.shared import RGBColor
    d.styles["Normal"].font.size = Pt(10)
    d.styles["Normal"].font.color.rgb = RGBColor(0x14, 0x14, 0x13)
    # the interface's look: serif headings, the title in ink, sections in terracotta
    for name, colour in (("Title", (0x14, 0x14, 0x13)), ("Heading 1", (0xC6, 0x61, 0x3F)),
                         ("Heading 2", (0x14, 0x14, 0x13))):
        style = d.styles[name]
        style.font.name = "Georgia"
        style.font.bold = False
        style.font.color.rgb = RGBColor(*colour)

    tokens = re.split(r"(<h1>.*?</h1>|<h2[^>]*>.*?</h2>|<h3>.*?</h3>|<table[^>]*>.*?</table>|"
                      r"<p[^>]*>.*?</p>|<li>.*?</li>)", html_text, flags=re.S)

    def strip(s):
        return html.unescape(re.sub(r"<[^>]+>", "", re.sub(r"<br\s*/?>", "\n", s))).strip()

    for tok in tokens:
        if tok.startswith("<h1>"):
            d.add_heading(strip(tok), level=0)
        elif tok.startswith("<h2"):
            d.add_heading(strip(tok), level=1)
        elif tok.startswith("<h3>"):
            d.add_heading(strip(tok), level=2)
        elif tok.startswith("<li>"):
            d.add_paragraph(strip(tok), style="List Bullet")
        elif tok.startswith("<table"):
            rows = re.findall(r"<tr>(.*?)</tr>", tok, flags=re.S)
            cells = [re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", r, flags=re.S) for r in rows]
            if cells:
                table = d.add_table(rows=len(cells), cols=len(cells[0]))
                table.style = "Light Grid Accent 1"
                for i, row in enumerate(cells):
                    for j, c in enumerate(row):
                        if j < len(table.columns):
                            table.cell(i, j).text = strip(c)
        elif tok.startswith("<p"):
            m = re.search(r"src='fig://([a-z_]+)'", tok)
            if m:
                d.add_picture(io.BytesIO(figures[m.group(1)]), width=Inches(6.3))
            else:
                text = strip(tok)
                if text:
                    d.add_paragraph(text)
    d.save(path)


def export_report(path: str, wall, lang: str, heights=None) -> None:
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        export_pdf(path, wall, lang, heights)
    elif ext == "docx":
        export_docx(path, wall, lang, heights)
    else:
        export_html(path if ext in ("html", "htm") else path + ".html", wall, lang, heights)
